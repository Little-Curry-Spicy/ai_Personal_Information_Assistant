#!/usr/bin/env python3
"""Generate a Chinese Word resume for Agent + fullstack interviews."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
ACCENT = RGBColor(0x2C, 0x5F, 0x8A)
BODY = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, *, size_pt: float, bold: bool = False, color: RGBColor = BODY, east_asia: str = "PingFang SC") -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_spacing(p, *, before=0, after=0, line=1.08, exact: int | None = None) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if exact is not None:
        pf.line_spacing = Pt(exact)
    else:
        pf.line_spacing = line
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def shade_paragraph(p, hex_color: str) -> None:
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    ppr.append(shd)


def add_bottom_border(p, color: str = "2C5F8A", sz: str = "12") -> None:
    ppr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    ppr.append(pBdr)


def add_text(p, text: str, **kwargs) -> None:
    run = p.add_run(text)
    set_run_font(run, **kwargs)


def heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=10, after=4, line=1.0)
    add_bottom_border(p)
    add_text(p, text, size_pt=12.5, bold=True, color=NAVY, east_asia="PingFang SC")


def job_header(doc, left: str, right: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=1, line=1.05)
    add_text(p, left, size_pt=11, bold=True, color=NAVY)
    add_text(p, "    ", size_pt=11, color=MUTED)
    add_text(p, right, size_pt=10, bold=False, color=MUTED)


def job_meta(doc, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=3, line=1.05)
    add_text(p, text, size_pt=9.5, color=ACCENT)


def bullet(doc, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=1, after=2, line=1.12)
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    add_text(p, "•  ", size_pt=10.5, color=ACCENT)
    if bold_lead:
        add_text(p, bold_lead, size_pt=10.5, bold=True, color=BODY)
        add_text(p, text, size_pt=10.5, color=BODY)
    else:
        add_text(p, text, size_pt=10.5, color=BODY)


def skill_line(doc, label: str, content: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=1, after=2, line=1.12)
    add_text(p, "•  ", size_pt=10.5, color=ACCENT)
    add_text(p, label, size_pt=10.5, bold=True, color=BODY)
    add_text(p, content, size_pt=10.5, color=BODY)


def build() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    # --- Name ---
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=2, line=1.0)
    add_text(name, "唐硕堃", size_pt=22, bold=True, color=NAVY, east_asia="PingFang SC")
    add_text(name, "  ·  Ben", size_pt=13, bold=False, color=ACCENT)

    intent = doc.add_paragraph()
    intent.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(intent, before=0, after=3, line=1.0)
    add_text(intent, "求职意向：AI Agent 开发工程师（全栈可兼）  ·  工作 7 年  ·  本科  ·  薪资面议", size_pt=10.5, color=BODY)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, before=0, after=6, line=1.0)
    add_bottom_border(contact, "1B3A5F", "18")
    add_text(contact, "18623806693  ·  tskwangyi@163.com  ·  GitHub: github.com/Little-Curry-Spicy", size_pt=9.5, color=MUTED)

    # --- Summary ---
    heading(doc, "个人简介")
    summary = doc.add_paragraph()
    set_paragraph_spacing(summary, before=2, after=2, line=1.15)
    add_text(
        summary,
        "7 年工程经验，前端出身、可独立闭环中小型全栈交付，近一年聚焦 Agent 与 RAG 系统。"
        "能用 LangGraph 编排带自纠正的多步工作流，结合混合检索、查询改写与评测门控提升问答可靠性；"
        "同时具备 React / Vue3 / TypeScript 与 Node.js（NestJS）接口、数据层和 Docker 交付能力。"
        "目标岗位以 Agent 为主，全栈作为落地与产品化补充。",
        size_pt=10.5,
        color=BODY,
    )

    # --- Skills ---
    heading(doc, "核心技能")
    skill_line(
        doc,
        "Agent / RAG：",
        "LangGraph 状态机与检查点、LangChain、Agentic RAG（检索→评估→改写→补检）、"
        "HyDE、查询改写、工具调用与最大重试；DeepSeek、DashScope Embedding；Chroma、BM25、RRF 混合检索。",
    )
    skill_line(
        doc,
        "服务端与数据：",
        "FastAPI + SSE 流式输出；Node.js / NestJS 模块化 API；PostgreSQL（PostgresSaver）、MySQL、Redis、Prisma；Docker 部署与联调。",
    )
    skill_line(
        doc,
        "前端与工程化：",
        "精通 React（3 年+）、Vue3（2 年+）、TypeScript；Next.js SSR；Vite / Webpack、Monorepo；性能治理（首屏 3.5s → 1.2s）。",
    )
    skill_line(
        doc,
        "质量与协作：",
        "Jest / Cypress；契约化联调（错误码、分页、状态机）；Hardhat 合约测试与多签流程协作（加分项，非主方向）。",
    )

    # --- Experience ---
    heading(doc, "工作经历")

    # 1. xone
    job_header(doc, "xone", "2025.05 — 2025.11  ·  北京")
    job_meta(doc, "全栈工程师  ·  Web3 钱包  ·  Next.js / NestJS / Solidity  ·  以太坊 + IPFS")
    bullet(
        doc,
        "从 0 到 1 交付 xone 钱包全栈：Next.js 账户 / 资产 / 交易页，NestJS 模块化 API 对齐链上余额、nonce、交易哈希与流水，打通转账—签名上链—到账确认闭环。",
        bold_lead="全栈闭环：",
    )
    bullet(
        doc,
        "基于链上持仓与交互行为，用 Agent 做交易风险提示和个性化资产 / DApp 建议，完成策略、工具调用与前后端联调，把 Agent 接到真实签名与转账场景。",
        bold_lead="Agent 助手：",
    )
    bullet(
        doc,
        "封装钱包连接、多链网络切换、交易签名、状态反馈与失败重试；用 Solidity 实现转账、授权与多签金库等核心合约，Hardhat 覆盖关键路径，敏感操作走多签提案与执行。",
        bold_lead="Web3 交互：",
    )
    bullet(
        doc,
        "批量交易与 Layer2 将用户 Gas 成本约降 40%；IPFS Gateway + CDN 加速 NFT / 资产元数据加载；首屏优化至约 2s，签名登录与交易确认体验贴近 Web2。",
        bold_lead="体验与成本：",
    )

    # 2. knowledge base
    job_header(doc, "企业知识库问答系统（Agentic RAG）", "2026.02 — 2026.06")
    job_meta(doc, "AI Agent 开发工程师（全栈）  ·  LangGraph / LangChain / FastAPI / Chroma")
    bullet(
        doc,
        "针对传统单轮 RAG 检索不稳、版本混淆、口语问法难匹配等问题，设计并落地带自纠正能力的 Agentic RAG，提升企业问答准确率与可运维性。",
        bold_lead="项目背景：",
    )
    bullet(
        doc,
        "用 LangGraph 状态机构建「基线检索 → 质量评估 → 查询改写 → 定向补检」迭代链路；设置最大重试避免死循环，评估通过后才进入答案生成；PostgresSaver 持久化 Agent 状态，支撑多轮会话恢复。",
        bold_lead="工作流编排：",
    )
    bullet(
        doc,
        "Chroma 向量检索 + BM25 关键词检索，经 RRF 融合；首轮轻量向量检索保速度，补检阶段自动切换混合检索扩召回。",
        bold_lead="多策略检索：",
    )
    bullet(
        doc,
        "底层术语映射零延迟打底，中层 LLM 做口语归一化并补全专业术语，顶层 HyDE 语义增强；按评估结果自动补齐时间约束与技术词，生成更贴近知识库的检索 Query，无需用户改写。",
        bold_lead="口语与改写：",
    )
    bullet(
        doc,
        "FastAPI 提供问答 API，SSE 流式返回思考 / 检索 / 生成过程。复杂场景（版本混淆、多维提问）检索召回 71% → 92%，整体问答准确率约提升 27%，显著减少人工调参与切片策略成本。",
        bold_lead="交付与结果：",
    )

    # 3. Dongchi
    job_header(doc, "浙江东驰科技有限公司", "2022.05 — 2025.05  ·  浙江")
    job_meta(doc, "高级全栈工程师  ·  浙江省国资委 SaaS 董事会管理系统（国家级优秀应用奖）")
    bullet(
        doc,
        "Vue3 + TypeScript 多租户架构，按租户配置动态加载组件、菜单与路由，一套应用服务多家国企，维护成本约降 70%。",
        bold_lead="前端架构：",
    )
    bullet(
        doc,
        "可视化流程引擎（拖拽表单设计器 + JSON Schema 驱动渲染），将流程定制周期由约 2 周压缩至约 2 天；落地 RBAC（菜单 / 按钮 / 数据域），自定义指令 v-permission + 路由守卫，前后端双重鉴权。",
        bold_lead="低代码与权限：",
    )
    bullet(
        doc,
        "对接租户隔离与治理类业务接口，统一错误码、分页与导出契约；排查跨租户串单与缓存不一致，保障会议、决议、报送等关键链路。系统获国家级优秀应用奖，用户满意度约 95%。",
        bold_lead="联调与成果：",
    )

    # 4. Cidi
    job_header(doc, "次第科技有限公司", "2020.06 — 2022.05  ·  郑州")
    job_meta(doc, "前端负责人  ·  数字资产 OTC 交易系统")
    bullet(
        doc,
        "以 TypeScript 统一模块边界；加密与大数运算下沉 WebAssembly，并结合 Web Worker，避免阻塞主线程。首屏 3.5s → 1.2s，核心加密模块约 5 倍性能提升。",
        bold_lead="性能：",
    )
    bullet(
        doc,
        "统一跨端框架一套代码输出 Web / Android / iOS，复用约 85%；推动规范、分支策略与自动化流水线，版本迭代周期约缩短 50%。",
        bold_lead="跨端与工程化：",
    )

    # --- Education ---
    heading(doc, "教育背景")
    edu = doc.add_paragraph()
    set_paragraph_spacing(edu, before=4, after=2, line=1.05)
    add_text(edu, "安阳工学院", size_pt=11, bold=True, color=NAVY)
    add_text(edu, "    自动化  ·  本科    2014.09 — 2018.06", size_pt=10.5, color=MUTED)

    # --- Self review ---
    heading(doc, "自我评价")
    review = doc.add_paragraph()
    set_paragraph_spacing(review, before=2, after=0, line=1.15)
    add_text(
        review,
        "擅长把 Agent 做成可约束的工程系统，而不是单次 Prompt：用状态机、评测门控和重试上限保证检索—生成链路可观测、可终止。"
        "全栈侧能独立完成 API、数据层与前端交互，把 Agent 能力接到真实业务（推荐、知识库问答、流式 UI）。"
        "工作风格重视验收标准与边界，能在 Agent 主路径上深挖，同时用全栈能力把方案推到可演示、可上线。",
        size_pt=10.5,
        color=BODY,
    )

    return doc


def main() -> None:
    doc = build()
    desktop = Path("/Users/tangshuokun/Desktop/唐硕堃-AI_Agent全栈工程师-简历.docx")
    project = Path("/Users/tangshuokun/学习/ai_Personal_Information_Assistant/唐硕堃-AI_Agent全栈工程师-简历.docx")
    for path in (desktop, project):
        doc.save(path)
        print(f"wrote {path}")


if __name__ == "__main__":
    main()
